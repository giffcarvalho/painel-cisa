"""Endpoints do Extrator de Dados."""

from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.database import get_db
from app.services.extrator_catalogo import FIELD_BY_ID, TIPOS_TABELA, listar_campos
from app.schemas.extrator_dados import (
    CatalogoResponse,
    ExportRequest,
    FiltroBuscaResponse,
    FiltrosResponse,
    PreviewRequest,
    PreviewResponse,
    TiposTabelaResponse,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None

router = APIRouter()
logger = logging.getLogger(__name__)

MAX_EXPORT_ROWS = 50000

AZUL_LABEL = "1F497D"
AZUL_TABELA = "1F4E78"
CINZA_TEXTO = "595959"
CINZA_BORDA = "D9E2EC"
BRANCO = "FFFFFF"

FILTROS_PERMITIDOS = {
    "municipio": {
        "sigla_uf": {
            "label": "UF",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "cod_municipio": {
            "label": "Município",
            "tipo": "integer",
            "cast": "int[]",
            "busca": True,
        },
        "regiao": {
            "label": "Região",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "semiarido_2022": {
            "label": "Semiárido",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "amazonia_legal": {
            "label": "Amazônia Legal",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "vale_jequetinhonha": {
            "label": "Vale do Jequitinhonha",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
    },
    "setor_censitario": {
        "sigla_uf": {
            "label": "UF",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "cod_municipio": {
            "label": "Município",
            "tipo": "integer",
            "cast": "int[]",
            "busca": True,
        },
        "regiao": {
            "label": "Região",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "semiarido_2022": {
            "label": "Semiárido",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "amazonia_legal": {
            "label": "Amazônia Legal",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "situacao": {
            "label": "Situação",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "situacao_detalhada": {
            "label": "Situação detalhada",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "tipo": {
            "label": "Tipo",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "com_dados": {
            "label": "Com dados",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "com_pessoas": {
            "label": "Com pessoas",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
        "todos_dados_omitidos": {
            "label": "Dados omitidos",
            "tipo": "boolean",
            "cast": "boolean[]",
            "busca": False,
        },
    },
    "instrumento": {
        "uf": {
            "label": "UF",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "situacao_contratacao": {
            "label": "Situação da contratação",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "fase_instrumento": {
            "label": "Fase do instrumento",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "acao_padronizada": {
            "label": "Ação",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "tipo_instrumento": {
            "label": "Tipo de instrumento",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "nome_proponente": {
            "label": "Proponente",
            "tipo": "text",
            "cast": "text[]",
            "busca": True,
        },
        "ano_proposta": {
            "label": "Ano da proposta",
            "tipo": "integer",
            "cast": "int[]",
            "busca": False,
        },
        "nr_proposta": {
            "label": "Número da proposta",
            "tipo": "text",
            "cast": "text[]",
            "busca": True,
            "sql": "nr_proposta::text",
        },
        "nr_instrumento": {
            "label": "Número do instrumento",
            "tipo": "text",
            "cast": "text[]",
            "busca": True,
            "sql": "nr_instrumento::text",
        },
        "operacao": {
            "label": "Operação",
            "tipo": "text",
            "cast": "text[]",
            "busca": True,
            "sql": "operacao::text",
        },
        "carteira_ativa": {
            "label": "Carteira ativa",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "novo_pac": {
            "label": "Novo PAC",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "componente": {
            "label": "Componente",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
        "coordenacao": {
            "label": "Coordenação",
            "tipo": "text",
            "cast": "text[]",
            "busca": False,
        },
    },
}

def _view(tipo_tabela: str) -> str:
    config = TIPOS_TABELA.get(tipo_tabela)
    if not config:
        raise HTTPException(status_code=400, detail="Tipo de tabela invalido.")
    return config["view"]

async def _execute_query(db: AsyncSession, sql: str, params: dict | None = None):
    try:
        return await db.execute(text(sql), params or {})
    except SQLAlchemyError as exc:
        logger.exception("Erro no banco de dados do Extrator: %s", exc)
        raise HTTPException(
            status_code=500,
            detail="Erro interno ao consultar a base de dados.",
        )

def _normalizar_lista(value: Any) -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [item for item in value if item not in (None, "")]
    return [value]

def _campos_solicitados(tipo_tabela: str, field_ids: list[str]) -> list[dict]:
    campos: list[dict] = []

    for field_id in field_ids:
        campo = FIELD_BY_ID.get(field_id)
        if not campo or campo["tipo_tabela"] != tipo_tabela or not campo["exportavel"]:
            raise HTTPException(
                status_code=400,
                detail=f"Campo nao permitido para esta tabela: {field_id}",
            )
        campos.append(campo)

    return campos

def _build_select(campos: list[dict]) -> str:
    return ",\n            ".join(
        f'{campo["column"]} AS "{campo["id"]}"'
        for campo in campos
    )

def _build_where(tipo_tabela: str, filtros: dict[str, Any]) -> tuple[str, dict]:
    permitidos = FILTROS_PERMITIDOS[tipo_tabela]
    clauses: list[str] = []
    params: dict[str, Any] = {}

    for filtro, raw_value in (filtros or {}).items():
        if filtro not in permitidos:
            raise HTTPException(
                status_code=400,
                detail=f"Filtro nao permitido para {tipo_tabela}: {filtro}",
            )

        values = _normalizar_lista(raw_value)
        if not values:
            continue

        config = permitidos[filtro]
        param_key = f"f_{filtro}"
        sql_column = config.get("sql", filtro)

        if config["tipo"] == "integer":
            try:
                params[param_key] = [int(v) for v in values]
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail=f"Filtro {filtro} deve ser numerico.")
        elif config["tipo"] == "boolean":
            params[param_key] = [str(v).lower() in {"true", "1", "sim"} for v in values]
        else:
            params[param_key] = [str(v).strip() for v in values if str(v).strip()]

        if tipo_tabela == "instrumento" and filtro == "uf":
            params[param_key] = [str(v).strip().upper() for v in values if str(v).strip()]
            clauses.append(
                f"""
                EXISTS (
                    SELECT 1
                    FROM regexp_split_to_table(COALESCE(uf::text, ''), '\s*,\s*') AS uf_item(valor)
                    WHERE upper(btrim(uf_item.valor)) = ANY(CAST(:f_uf AS text[]))
                )
                """
            )
            continue

        clauses.append(f"{sql_column} = ANY(CAST(:{param_key} AS {config['cast']}))")

    return ("WHERE " + " AND ".join(clauses), params) if clauses else ("", params)

def _order_by(tipo_tabela: str) -> str:
    columns = TIPOS_TABELA[tipo_tabela]["ordem"]
    return ", ".join(columns)

async def _count_rows(db: AsyncSession, tipo_tabela: str, where: str, params: dict) -> int:
    result = await _execute_query(
        db,
        f"""
        SELECT COUNT(*)
        FROM {_view(tipo_tabela)}
        {where}
        """,
        params,
    )
    return int(result.scalar_one() or 0)

async def _metadata(db: AsyncSession, tipo_tabela: str, where: str, params: dict) -> dict[str, Any]:
    if tipo_tabela != "instrumento":
        return {}

    result = await _execute_query(
        db,
        f"""
        SELECT
            MAX(data_dados_transferegov) AS data_dados_transferegov,
            MAX(data_dados_caixa) AS data_dados_caixa
        FROM {_view(tipo_tabela)}
        {where}
        """,
        params,
    )
    row = dict(result.mappings().one())
    return jsonable_encoder(row)

def _formatar_data_br(valor: Any) -> str:
    if valor is None:
        return "Não disponível"

    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")

    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    if isinstance(valor, str):
        valor_normalizado = valor.strip()

        if not valor_normalizado:
            return "Não disponível"

        try:
            if "T" in valor_normalizado or " " in valor_normalizado:
                data_convertida = datetime.fromisoformat(
                    valor_normalizado.replace("Z", "+00:00")
                )
                return data_convertida.strftime("%d/%m/%Y")

            data_convertida = date.fromisoformat(valor_normalizado)
            return data_convertida.strftime("%d/%m/%Y")
        except ValueError:
            return "Não disponível"

    return "Não disponível"

def _formatar_filtros(tipo_tabela: str, filtros: dict[str, Any]) -> str:
    if not filtros:
        return "Sem filtros aplicados"

    permitidos = FILTROS_PERMITIDOS.get(tipo_tabela, {})
    partes = []

    for key, value in filtros.items():
        values = _normalizar_lista(value)
        if not values:
            continue

        label = permitidos.get(key, {}).get("label", key)
        valores_formatados = []

        for item in values:
            if isinstance(item, bool):
                valores_formatados.append("Sim" if item else "Não")
            elif str(item).lower() == "true":
                valores_formatados.append("Sim")
            elif str(item).lower() == "false":
                valores_formatados.append("Não")
            else:
                valores_formatados.append(str(item))

        partes.append(f"{label}: {', '.join(valores_formatados)}")

    return "; ".join(partes) if partes else "Sem filtros aplicados"

@router.get("/tipos-tabela", response_model=TiposTabelaResponse)
async def get_tipos_tabela(response: Response):
    response.headers["Cache-Control"] = "no-cache"
    return {
        "data": [
            {"id": key, "label": value["label"], "descricao": value["descricao"]}
            for key, value in TIPOS_TABELA.items()
        ]
    }

@router.get("/catalogo", response_model=CatalogoResponse)
async def get_catalogo(tipo_tabela: str, response: Response):
    if tipo_tabela not in TIPOS_TABELA:
        raise HTTPException(status_code=400, detail="Tipo de tabela invalido.")
    response.headers["Cache-Control"] = "public, max-age=3600"
    return {"tipo_tabela": tipo_tabela, "data": listar_campos(tipo_tabela)}

@router.get("/filtros", response_model=FiltrosResponse)
async def get_filtros(tipo_tabela: str, response: Response, db: AsyncSession = Depends(get_db)):
    if tipo_tabela not in TIPOS_TABELA:
        raise HTTPException(status_code=400, detail="Tipo de tabela invalido.")

    response.headers["Cache-Control"] = "public, max-age=1800"

    filtros_config = FILTROS_PERMITIDOS[tipo_tabela]
    select_parts = []
    for campo, config in filtros_config.items():
        if config.get("busca"):
            continue

        if tipo_tabela == "instrumento" and campo == "uf":
            select_parts.append(
                "(SELECT array_agg(sigla_uf ORDER BY sigla_uf) FROM territorio.tb_uf) AS uf"
            )
            continue

        sql_column = config.get("sql", campo)
        select_parts.append(
            f"array_remove(array_agg(DISTINCT {sql_column} ORDER BY {sql_column}), NULL) AS {campo}"
        )

    if not select_parts:
        return {"tipo_tabela": tipo_tabela, "data": []}

    result = await _execute_query(
        db,
        f"""
        SELECT
            {", ".join(select_parts)}
        FROM {_view(tipo_tabela)}
        """,
    )
    row = dict(result.mappings().one())

    return {
        "tipo_tabela": tipo_tabela,
        "data": [
            {
                "campo": campo,
                "label": config["label"],
                "tipo_dado": config["tipo"],
                "opcoes": row.get(campo) or [],
                "busca": config.get("busca", False),
            }
            for campo, config in filtros_config.items()
        ],
    }

@router.get("/filtros/busca", response_model=FiltroBuscaResponse)
async def buscar_filtro(
    tipo_tabela: str,
    campo: str,
    response: Response,
    q: str = Query("", max_length=100),
    limit: int = Query(50, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
):
    if tipo_tabela not in TIPOS_TABELA:
        raise HTTPException(status_code=400, detail="Tipo de tabela invalido.")

    config = FILTROS_PERMITIDOS[tipo_tabela].get(campo)
    if not config or not config.get("busca"):
        raise HTTPException(status_code=400, detail="Filtro de busca nao permitido.")

    response.headers["Cache-Control"] = "private, max-age=120"

    value_expr = config.get("sql", campo)

    if campo == "cod_municipio":
        nome_col = "nome" if tipo_tabela == "municipio" else "nome_municipio"

        if tipo_tabela == "municipio":
            label_expr = nome_col
        else:
            label_expr = f"({nome_col} || ' - ' || sigla_uf)"

        search_expr = f"(CAST({value_expr} AS text) || ' ' || COALESCE({nome_col}, '') || ' ' || COALESCE(sigla_uf, ''))"
    else:
        label_expr = f"CAST({value_expr} AS text)"
        search_expr = f"CAST({value_expr} AS text)"

    termo = q.strip()

    result = await _execute_query(
        db,
        f"""
        SELECT DISTINCT
            CAST({value_expr} AS text) AS value,
            {label_expr} AS label
        FROM {_view(tipo_tabela)}
        WHERE {value_expr} IS NOT NULL
          AND (:q = '' OR {search_expr} ILIKE :q_like)
        ORDER BY label
        LIMIT :limit
        """,
        {
            "q": termo,
            "q_like": f"%{termo}%",
            "limit": limit,
        },
    )

    return {
        "tipo_tabela": tipo_tabela,
        "campo": campo,
        "data": [dict(row) for row in result.mappings().all()],
    }

@router.post("/previa", response_model=PreviewResponse)
async def post_previa(payload: PreviewRequest, db: AsyncSession = Depends(get_db)):
    campos = _campos_solicitados(payload.tipo_tabela, payload.field_ids)
    campos = sorted(
        campos,
        key=lambda campo: 0 if campo["papel_semantico"] == "dimensao" else 1,
    )
    select_sql = _build_select(campos)
    where, params = _build_where(payload.tipo_tabela, payload.filtros)

    total = await _count_rows(db, payload.tipo_tabela, where, params)
    metadados = await _metadata(db, payload.tipo_tabela, where, params)

    data_params = {**params, "limit": payload.limit}

    result = await _execute_query(
        db,
        f"""
        SELECT
            {select_sql}
        FROM {_view(payload.tipo_tabela)}
        {where}
        ORDER BY {_order_by(payload.tipo_tabela)}
        LIMIT :limit
        """,
        data_params,
    )

    return {
        "tipo_tabela": payload.tipo_tabela,
        "total_estimado": total,
        "limit": payload.limit,
        "columns": campos,
        "data": jsonable_encoder([dict(row) for row in result.mappings().all()]),
        "metadados": metadados,
    }

@router.post("/exportar/excel")
async def post_exportar_excel(payload: ExportRequest, db: AsyncSession = Depends(get_db)):
    if Workbook is None:
        raise HTTPException(
            status_code=500,
            detail="A dependencia openpyxl nao esta instalada no backend.",
        )

    campos = _campos_solicitados(payload.tipo_tabela, payload.field_ids)
    select_sql = _build_select(campos)
    where, params = _build_where(payload.tipo_tabela, payload.filtros)

    total = await _count_rows(db, payload.tipo_tabela, where, params)
    if total > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=413,
            detail=f"A exportacao possui {total} linhas. Refine os filtros ou limite a extracao a ate {MAX_EXPORT_ROWS} linhas.",
        )

    metadados = await _metadata(db, payload.tipo_tabela, where, params)

    result = await _execute_query(
        db,
        f"""
        SELECT
            {select_sql}
        FROM {_view(payload.tipo_tabela)}
        {where}
        ORDER BY {_order_by(payload.tipo_tabela)}
        """,
        params,
    )
    rows = [dict(row) for row in result.mappings().all()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Extrator"
    ws.sheet_view.showGridLines = False

    tipo_label = TIPOS_TABELA[payload.tipo_tabela]["label"]
    data_exportacao = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    borda_leve = Border(
        left=Side(style="thin", color=CINZA_BORDA),
        right=Side(style="thin", color=CINZA_BORDA),
        top=Side(style="thin", color=CINZA_BORDA),
        bottom=Side(style="thin", color=CINZA_BORDA),
    )

    borda_total = Border(
        left=Side(style="thin", color=CINZA_BORDA),
        right=Side(style="thin", color=CINZA_BORDA),
        top=Side(style="medium", color=AZUL_TABELA),
        bottom=Side(style="thin", color=CINZA_BORDA),
    )

    alinhamento_metadado = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    alinhamento_cabecalho = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    alinhamento_dados = Alignment(
        vertical="center",
    )

    # Cabeçalho fixo: linhas 1 e 2.
    ws.merge_cells("A1:B1")
    ws["A1"] = "Painel DSR"
    ws["A1"].font = Font(
        name="Cambria",
        size=14,
        bold=True,
    )
    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20

    filtros_aplicados = _formatar_filtros(
        payload.tipo_tabela,
        payload.filtros,
    )
    data_transferegov = _formatar_data_br(
        metadados.get("data_dados_transferegov")
    )
    data_caixa = _formatar_data_br(
        metadados.get("data_dados_caixa")
    )

    metadados_layout = [
        ("A2", "B2", "Filtros aplicados", filtros_aplicados),
        ("C1", "D1", "Tabela", tipo_label),
        ("C2", "D2", "Data e hora da exportação", data_exportacao),
        ("E1", "F1", "Data dados Transferegov", data_transferegov),
        ("E2", "F2", "Data dados Caixa", data_caixa),
    ]

    larguras_metadados: dict[str, list[str]] = {}

    for label_ref, value_ref, label, value in metadados_layout:
        label_cell = ws[label_ref]
        value_cell = ws[value_ref]

        label_cell.value = label
        label_cell.font = Font(
            name="Calibri",
            size=10,
            color=AZUL_LABEL,
            bold=True,
        )
        label_cell.alignment = alinhamento_metadado

        value_cell.value = value
        value_cell.font = Font(
            name="Calibri",
            size=10,
            color=CINZA_TEXTO,
            bold=False,
        )
        value_cell.alignment = alinhamento_metadado

        larguras_metadados.setdefault(label_cell.column_letter, []).append(label)
        larguras_metadados.setdefault(value_cell.column_letter, []).append(
            str(value or "")
        )

    # A tabela começa obrigatoriamente na linha 3.
    table_header_row = 3

    for idx, campo in enumerate(campos, start=1):
        cell = ws.cell(
            row=table_header_row,
            column=idx,
            value=campo["label"],
        )
        cell.fill = PatternFill("solid", fgColor=AZUL_TABELA)
        cell.font = Font(
            name="Calibri",
            size=10,
            color=BRANCO,
            bold=True,
        )
        cell.alignment = alinhamento_cabecalho
        cell.border = borda_leve

    ws.row_dimensions[table_header_row].height = 34

    # Os dados começam obrigatoriamente na linha 4.
    for row in rows:
        ws.append([row.get(campo["id"]) for campo in campos])

        for cell in ws[ws.max_row]:
            cell.alignment = alinhamento_dados
            cell.border = borda_leve

    # Intervalo dos dados: é definido antes da criação de qualquer total.
    data_first_row = table_header_row + 1
    data_last_row = table_header_row + len(rows)

    # Apenas métricas financeiras/monetárias podem ser totalizadas.
    # Percentuais, quantidades, códigos, rankings, datas e demais métricas ficam fora.
    colunas_totalizaveis = [
        idx
        for idx, campo in enumerate(campos, start=1)
        if (
            campo["papel_semantico"] == "metrica"
            and campo["tipo_dado"] == "currency"
            and campo["formato"] == "brl"
        )
    ]

    # Mantém ajuste automático de largura e formatos numéricos existentes.
    for idx, campo in enumerate(campos, start=1):
        valores_amostra = [campo["label"]] + [
            str(row.get(campo["id"]) or "")
            for row in rows[:200]
        ]
        largura = min(
            max(max(len(valor) for valor in valores_amostra) + 2, 14),
            45,
        )

        ws.column_dimensions[get_column_letter(idx)].width = largura

        if campo["formato"] == "brl":
            for cells in ws.iter_rows(
                min_row=data_first_row,
                max_row=data_last_row,
                min_col=idx,
                max_col=idx,
            ):
                cells[0].number_format = "R$ #,##0.00"
        elif campo["formato"] == "percent":
            for cells in ws.iter_rows(
                min_row=data_first_row,
                max_row=data_last_row,
                min_col=idx,
                max_col=idx,
            ):
                cells[0].number_format = "0.00%"

    # Esta seção fica fora do laço de formatação: cria uma única linha final.
    if colunas_totalizaveis:
        total_row = data_last_row + 1

        # Usa a primeira coluna não financeira para o rótulo — normalmente a coluna A.
        coluna_rotulo = next(
            (
                idx
                for idx in range(1, len(campos) + 1)
                if idx not in colunas_totalizaveis
            ),
            None,
        )

        # Caso excepcional: exportação composta somente por colunas financeiras.
        # O primeiro total mostra visualmente "Total geral — R$ ...", sem transformar
        # a célula em texto e sem perder a soma dessa coluna.
        rotulo_embutido_no_total = coluna_rotulo is None
        coluna_rotulo = coluna_rotulo or colunas_totalizaveis[0]

        for idx in range(1, len(campos) + 1):
            cell = ws.cell(row=total_row, column=idx)
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.alignment = alinhamento_dados
            cell.border = borda_total

            if idx in colunas_totalizaveis:
                column_letter = get_column_letter(idx)
                cell.value = (
                    f"=SUBTOTAL(109,{column_letter}{data_first_row}:{column_letter}{data_last_row})"
                    if rows
                    else 0
                )
                cell.number_format = (
                    '"Total geral — "R$ #,##0.00'
                    if rotulo_embutido_no_total and idx == coluna_rotulo
                    else "R$ #,##0.00"
                )
            elif idx == coluna_rotulo:
                cell.value = "Total geral"

    # Garante espaço para os metadados sem reduzir larguras da tabela.
    for column, valores in larguras_metadados.items():
        largura = min(
            max(max(len(valor) for valor in valores) + 2, 14),
            45,
        )
        ws.column_dimensions[column].width = max(
            ws.column_dimensions[column].width or 0,
            largura,
        )

    last_column = get_column_letter(len(campos))
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{last_column}{data_last_row}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"extrator_dados_{payload.tipo_tabela}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )